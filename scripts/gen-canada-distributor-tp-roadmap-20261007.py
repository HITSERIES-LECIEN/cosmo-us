#!/usr/bin/env python3
"""Generate Canada distributor sales roadmap + TP strategy Word doc (2026-10-07)."""

from datetime import date
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


OUT = Path(
    "/Users/hiroki/Documents/cosmo-us/docs/"
    "canada-distributor-direct-ship-roadmap-tp-strategy-20261007.docx"
)


def set_run_font(run, size=11, bold=False, color=None):
    run.font.name = "Yu Gothic"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Yu Gothic")
    run.font.size = Pt(size)
    run.bold = bold
    if color:
        run.font.color.rgb = color


def add_para(doc, text, *, size=11, bold=False, space_after=8, space_before=0):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(space_after)
    p.paragraph_format.space_before = Pt(space_before)
    p.paragraph_format.line_spacing = 1.35
    run = p.add_run(text)
    set_run_font(run, size=size, bold=bold)
    return p


def add_h(doc, text, level=1):
    h = doc.add_heading(text, level=level)
    for run in h.runs:
        set_run_font(run, size=16 if level == 1 else 13 if level == 2 else 11, bold=True)
    return h


def add_bullets(doc, items, *, size=11):
    for item in items:
        p = doc.add_paragraph(style="List Bullet")
        p.paragraph_format.space_after = Pt(4)
        p.paragraph_format.line_spacing = 1.3
        run = p.add_run(item)
        set_run_font(run, size=size)


def add_numbered(doc, items, *, size=11):
    for item in items:
        p = doc.add_paragraph(style="List Number")
        p.paragraph_format.space_after = Pt(4)
        p.paragraph_format.line_spacing = 1.3
        run = p.add_run(item)
        set_run_font(run, size=size)


def set_cell_text(cell, text, *, bold=False, size=9):
    cell.text = ""
    p = cell.paragraphs[0]
    run = p.add_run(text)
    set_run_font(run, size=size, bold=bold)


def add_table(doc, headers, rows, col_widths=None):
    table = doc.add_table(rows=1 + len(rows), cols=len(headers))
    table.style = "Table Grid"
    for i, h in enumerate(headers):
        set_cell_text(table.rows[0].cells[i], h, bold=True, size=9)
    for r_i, row in enumerate(rows):
        for c_i, val in enumerate(row):
            set_cell_text(table.rows[r_i + 1].cells[c_i], val, size=9)
    if col_widths:
        for row in table.rows:
            for i, w in enumerate(col_widths):
                row.cells[i].width = Inches(w)
    doc.add_paragraph()
    return table


def build():
    doc = Document()
    section = doc.sections[0]
    section.top_margin = Inches(0.85)
    section.bottom_margin = Inches(0.85)
    section.left_margin = Inches(0.95)
    section.right_margin = Inches(0.95)

    # --- Title ---
    t = doc.add_paragraph()
    t.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = t.add_run(
        "Canada Distributor 直送販売\nロードマップ及び移転価格（TP）戦略"
    )
    set_run_font(r, size=18, bold=True)

    meta = [
        f"作成日: {date(2026, 10, 7).isoformat()}",
        "対象: LECIEN Japan（LCJ）× TANAAKK INC × Canadian Distributor",
        "位置づけ: 社内方針メモ（法務・税務・通関の確定助言ではない）",
        "前提: 米経由による関税二重負担を避け、日本→カナダ直送でB2B卸販売を行う",
    ]
    for m in meta:
        p = add_para(doc, m, size=10, space_after=2)
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER

    add_para(
        doc,
        "本資料は、Canada distributorからの打診を受けた直送スキームについて、"
        "商業構造・契約・オペ・通関・GST/PE・移転価格を一連のロードマップとして整理したものである。"
        "米署名済み Master（Consignment + U.S. Principal）及び Amendment No.2 をそのままCanadaへ流用しないことを前提とする。",
        size=10,
        space_before=12,
    )

    # --- 1 Exec summary ---
    add_h(doc, "1. エグゼクティブサマリー", 1)

    add_h(doc, "1.1 推奨構造（第一候補）", 2)
    add_bullets(
        doc,
        [
            "SOR / 契約売主: TANAAKK INC（Canadian Distributorとの契約主体）",
            "IOR / 通関・輸入GST・関税: Canadian Distributor",
            "Brand / 日本在庫 owner / Japan fulfillment: LECIEN Japan（LCJ）",
            "物流: 日本 → カナダ直送（米を経由しない）",
            "上流: Consignment（米Path Bの自動流用は禁止。Canada Scheduleで法的性質を再定義）",
            "LCJはCanadian顧客との販売条件交渉・契約締結を行わない",
        ],
    )

    add_h(doc, "1.2 移転価格の現時点の結論", 2)
    add_bullets(
        doc,
        [
            "米刺しゅう15%、Gross 10%固定の機械流用は不採用",
            "外部commissionデータだけでは率を確定できない",
            "pure manufacturers’ repの参考帯は概ね Net Sales の 5–9%（平均付近7%、データは古く品目混在）",
            "TANAAKKはSOR・pricing・invoice・collection・与信等でpure repより重いため、7%超も理論上あり得る（定性）",
            "10%は「結論」ではなく「ROS仮説を満たす場合の候補（テスト仮説）」",
            "最終決定は TNMM（CanadaセグメントのEBIT/Sales ≈ 2–3.5%を目安）で行う",
            "実務の感度候補: 5% / 7.5% / 10% / 12.5%（採用はOPEX実績とROSレンジで決める）",
        ],
    )

    add_h(doc, "1.3 Go条件（要約）", 2)
    add_numbered(
        doc,
        [
            "Canada専用 Schedule / Amendment の起草・署名",
            "Canadian customs broker による三者構造（Vendor=TANAAKK / Shipper=LCJ / IOR=Distributor）の書面確認",
            "Incoterms・title移転・未収・不良・精算のCanada固有定義",
            "Canada関連OPEX配賦ルール確定と retention×ROS 感度",
            "加国GST/HST carrying on business 及びPEの税理士確認（規模シナリオ付き）",
        ],
    )

    # --- 2 Background ---
    add_h(doc, "2. 背景と目的", 1)
    add_para(
        doc,
        "Canadian DistributorからTANAAKK INCへ打診があった。"
        "米国在庫・米輸入を介すると米側関税等の後にカナダ再輸入コストが重なり、二重負担となる。"
        "このため日本からの直送を希望している。目的は、(1) 関税の二重負担回避、"
        "(2) North America商業の一貫性（顧客接点はTANAAKK）、"
        "(3) LCJのCanada PE/GST負荷の抑制、(4) アームズレングスなグループ間対価の設計、である。",
    )

    add_h(doc, "2.1 米モデルとの関係", 2)
    add_table(
        doc,
        ["項目", "米国（署名済みMaster）", "Canada（推奨）"],
        [
            ["契約シェル", "Consignment + U.S. Principal", "Canada専用Scheduleで再定義"],
            ["SOR", "TANAAKK", "TANAAKK"],
            ["IOR", "TANAAKK", "Canadian Distributor"],
            ["在庫（販売前）", "LCJ（3PL等）", "LCJ（日本）。CA常駐在庫なし"],
            ["直送改正", "Amendment No.2（Japan→U.S.）", "流用禁止。Canada条項を新設"],
            ["報酬", "刺しゅう Net 15/85 等", "Canada FARに基づく別率（15%流用禁止）"],
        ],
        col_widths=[1.3, 2.3, 2.6],
    )

    add_para(
        doc,
        "批判的注記: 「SORであること」と「LCJ→TANAAKKが通常のintercompany buy-sellであること」は同義ではない。"
        "米契約は意図的に分離している。Canadaでも真正な顧客向けcommercial saleがあれば、"
        "上流を直ちにbuy-sell必須とはしない。ただし米Path Bを沈黙のまま輸入せず、Canadaでtitle/精算を定義する。",
        size=10,
    )

    # --- 3 Structure ---
    add_h(doc, "3. 推奨商業・物流構造", 1)

    add_h(doc, "3.1 フロー", 2)
    add_para(doc, "Commercial Flow", bold=True, space_after=4)
    add_para(
        doc,
        "LECIEN（Consignment / inventory owner）→ TANAAKK INC（Sale for export to Canada / SOR）→ Canadian Distributor（Buyer）",
        size=10,
    )
    add_para(doc, "Physical Flow", bold=True, space_after=4)
    add_para(
        doc,
        "LECIEN Japan ────（直送）────→ Canadian Distributor",
        size=10,
    )

    add_h(doc, "3.2 役割（FAR要約）", 2)
    add_table(
        doc,
        ["主体", "機能・資産", "主なリスク", "持たないもの"],
        [
            [
                "LECIEN Japan",
                "Brand/IP、設計・調達・製造、日本在庫、Japan fulfillment（TANAAKK指図）",
                "製造・品質・出荷前在庫、瑕疵（契約上のindemnityに従う）",
                "CA顧客との交渉・契約・回収",
            ],
            [
                "TANAAKK INC",
                "Distributor関係、pricing、交渉、受注、与信、請求、回収、出荷指図、市場開発",
                "自社販管・MD費用、Distributor売掛（回収ベース精算と整合）",
                "CA在庫、原則Canadian IOR",
            ],
            [
                "Canadian Distributor",
                "購入、IOR、通関、国内流通、下流価格",
                "輸入関税・輸入GST、CA在庫、下流与信・流通",
                "グループ内principal機能",
            ],
        ],
        col_widths=[1.2, 2.2, 1.8, 1.5],
    )

    add_h(doc, "3.3 呼称上の注意", 2)
    add_para(
        doc,
        "対外契約では「Seller of Record / contracting seller」で足りる。"
        "TP文書で「Canadian Commercial Principal」と書く場合は、"
        "pre-sale inventory・製造リスクを負わない limited-risk / flash-title の商業売主であることを併記する。"
        "呼称と残余利益（LCJ）・routine報酬（TANAAKK）の不一致を避ける。",
    )

    # --- 4 Legal roadmap ---
    add_h(doc, "4. 契約・法務ロードマップ", 1)

    add_h(doc, "4.1 Canada Schedule / Amendment に入れる条項", 2)
    add_numbered(
        doc,
        [
            "Territory: Canada（U.S. Amendment No.2の適用除外を明示）",
            "TANAAKK = sole contracting seller / SOR vis-à-vis Canadian customers",
            "LCJはCanadian顧客とcommercial termsを交渉・決定・承諾しない",
            "LCJはTANAAKKのShipment Instructionに基づきJapanからCanadaへfulfillment",
            "IOR = Canadian Distributor。TANAAKKは別途書面がない限りCanadian IORにならない（DDP禁止を原則）",
            "Incoterms と title / risk 移転時点（例: FCA Japan 又は Schedule定義のtender時等）",
            "上流title移転の法的性質（consignment履行か、Canada用の独自定義か）。米Path Bの自動適用禁止",
            "返品・refund決定権（TANAAKK）と経済負担分界（LCJ/TANAAKK）",
            "未収・貸倒: TANAAKKが与信リスクを負う場合、LCJ精算は回収ベース又はbad debt控除後",
            "Net Sales定義、retention率又はレンジ、OM/ROS調整条項",
            "保険、クレーム、リコール、知財表示",
            "Customs documentation協力義務（CI・SI・BLのストーリー一致）",
        ],
    )

    add_h(doc, "4.2 Distributor契約（TANAAKK ↔ Canadian Distributor）", 2)
    add_bullets(
        doc,
        [
            "売買契約の当事者はTANAAKKとDistributorのみ（LCJは当事者にしない）",
            "支払先・請求書発行者 = TANAAKK",
            "IOR・通関・輸入税負担 = Distributor",
            "価格・支払条件・与信限度・季節条件",
            "再販価格はDistributorが決定（ブランドMAPがある場合は別紙政策として整理）",
            "保証・不良の窓口はTANAAKK（バックツーバックでLCJへ）",
        ],
    )

    add_h(doc, "4.3 フェーズ（法務）", 2)
    add_table(
        doc,
        ["Phase", "成果物", "完了条件"],
        [
            ["L0", "方針合意（本メモ級）", "SOR/IOR/直送/Consignment維持の方針承認"],
            ["L1", "Canada Schedule草稿", "title・精算・IOR・報酬仮説を条文化"],
            ["L2", "Distributor売買契約草稿", "Incoterms・IOR・支払と一致"],
            ["L3", "LCJ–TANAAKK署名", "Schedule発効"],
            ["L4", "Distributor署名・最初のPO", "Go-live"],
        ],
        col_widths=[0.8, 2.4, 3.0],
    )

    # --- 5 Ops customs ---
    add_h(doc, "5. オペ・通関・書類ロードマップ", 1)

    add_h(doc, "5.1 一貫させるべきストーリー", 2)
    add_bullets(
        doc,
        [
            "Canadian DistributorのPO → TANAAKK宛",
            "TANAAKKのSales Invoice → Distributor（Vendor/Seller = TANAAKK）",
            "TANAAKK → LCJ のShipment Instruction",
            "LCJのshipping documents（shipper = LCJ でも可）",
            "Canada customs declaration（IOR = Distributor、建値は原則TANAAKK請求額）",
            "Payment: Distributor → TANAAKK、グループ内精算: TANAAKK ⇄ LCJ（Scheduleどおり）",
        ],
    )
    add_para(
        doc,
        "BL/AWB shipper = LCJ、Commercial seller = TANAAKK、Buyer/IOR = Distributor は三角・drop shipmentとして矛盾ではない。"
        "重要なのはsellerとshipperの法人一致ではなく、上記ストーリーの一貫性である。",
        size=10,
    )

    add_h(doc, "5.2 Customs valuation", 2)
    add_para(
        doc,
        "カナダの課税価格は、通常、Canada向け輸出販売において購入者（Distributor）が支払う価格である。"
        "したがって LCJ→TANAAKK を形式的に $60、$80 の二段buy-sellにしても、"
        "単純に $60 をCanada申告額にできるわけではない。"
        "customs安全のためのbuy-sell必須化は必然ではない。"
        "再構成リスク（真の売主はLCJでは？）への防御は、商業分離と書類一貫、broker指示である。",
    )

    add_h(doc, "5.3 Broker事前確認（必須）", 2)
    add_para(doc, "以下をCanadian customs brokerへ提示し、書面回答を取得する。", space_after=6)
    add_bullets(
        doc,
        [
            "Vendor/Seller: TANAAKK INC",
            "Shipper: LECIEN（Japan）",
            "Importer of Record: Canadian Distributor",
            "Physical shipment: Japan → Canada",
            "Commercial Invoiceの記載項目、related-party開示の要否、valuation method",
        ],
    )

    add_h(doc, "5.4 オペフェーズ", 2)
    add_table(
        doc,
        ["Phase", "内容"],
        [
            ["O1", "Incoterms・梱包・HS・原産地・ラベルの商品別チェックリスト"],
            ["O2", "Shipment Instructionテンプレート（TANAAKK→LCJ）"],
            ["O3", "CI/PL/BLの見本セットをbrokerレビュー"],
            ["O4", "パイロット出荷1件（少額）で端から端まで検証"],
            ["O5", "定常SOP（PO→承認→指図→出荷→通関→入金→月次精算）"],
        ],
        col_widths=[0.8, 5.4],
    )

    # --- 6 Tax ---
    add_h(doc, "6. カナダ間接税・PE・日本側の留意", 1)

    add_h(doc, "6.1 GST/HST", 2)
    add_bullets(
        doc,
        [
            "Buyer=IORの場合、輸入GST（Division III）は原則Distributor負担",
            "place of supply と carrying on business は別問題。office/在庫なしでも登録要否があり得る",
            "単発〜少数B2Bで契約・承諾・請求が米側完結ならリスクは相対的に低いが、反復・主力化で再判定",
            "継続的モニタリングをSchedule又は社内ポリシーに明記",
        ],
    )

    add_h(doc, "6.2 PE", 2)
    add_bullets(
        doc,
        [
            "CAにconsignment在庫を置かない想定なら、在庫PEは米3PLモデルより軽い",
            "LCJの依存代理人化（TANAAKKがLCJのために契約慣行）を避けるため、SOR=TANAAKKの実態を維持",
            "日加・米加の条約PEは税理士確認",
        ],
    )

    add_h(doc, "6.3 日本側", 2)
    add_bullets(
        doc,
        [
            "輸出免税・船積書類と、consignment／title定義の会計処理を一致",
            "LCJ残余・TANAAKK routine のTP方針を日本側文書にも反映",
        ],
    )

    # --- 7 TP strategy ---
    add_h(doc, "7. 移転価格戦略（詳細）", 1)

    add_h(doc, "7.1 基本方針", 2)
    add_para(
        doc,
        "Canadaレグは、LCJ = product/brand/inventory principal（残余）、"
        "TANAAKK = limited-risk commercial seller / SOR（routine）、"
        "Distributor = 独立第三者、の三層とする。"
        "対価の主形式は Net Sales × retention%（又は同等の精算）。"
        "ALの本丸は retention% の表面値ではなく、TANAAKK Canadaセグメントの利益水準（OM/ROS）である。",
    )

    add_h(doc, "7.2 Tested party と PLI", 2)
    add_table(
        doc,
        ["項目", "内容"],
        [
            ["Tested party", "TANAAKK INC の Canada distributor セグメント"],
            ["主PLI", "EBIT / Canada Net Sales（Operating Margin）"],
            ["補足PLI", "必要に応じ Berry ratio（EBIT / OpEx）"],
            ["目標レンジ（provisional）", "OM 約 2.0–3.5%（中心イメージ 2–3%）"],
            ["契約率の位置づけ", "OMを達成するための導出値・調整可能なバンド"],
        ],
        col_widths=[2.0, 4.2],
    )

    add_h(doc, "7.3 比較可能性（何ができて何ができないか）", 2)
    add_para(doc, "できないこと", bold=True, space_after=4)
    add_bullets(
        doc,
        [
            "「JP consignment × 米SOR × CA distributor=IOR」と完全同型の競合％の横並び",
            "競合卸の消費者価格・卸粗利を、TANAAKK retentionの直接CUPにすること",
            "求人サイトのcommission掲載を主comparableにすること",
        ],
    )
    add_para(doc, "できること（本筋）", bold=True, space_after=4)
    add_bullets(
        doc,
        [
            "NAICS 424310（Piece Goods, Notions, and Other Dry Goods Merchant Wholesalers）を探索母集団の入口にする（thread / notions / embroidery products 等を含む）",
            "在庫の薄い販売会社・LRD・sales entity のOMをTNMMで比較",
            "manufacturers’ rep commissionを補足の相場感として用いる（主検定にしない）",
        ],
    )

    add_h(doc, "7.4 外部データの批判的整理", 2)
    add_table(
        doc,
        ["データ", "示唆", "限界・使い方"],
        [
            [
                "NAICS 424310",
                "業種近さは高い（探索universe）",
                "率・利益を直接決めない",
            ],
            [
                "Merchant wholesaler GM ≈37%",
                "フル卸の粗利水準",
                "TANAAKK retentionと直接比較禁止。除外説明のみ",
            ],
            [
                "MANA distributor向け佣金 5–9%（平均≈7%）",
                "pure repのballparkとして相対的に有用",
                "1999–2005調査。standardではない。品目混在。Apparel等は≈5%の観測も。repは通常invoice/与信なし",
            ],
            [
                "求人等のtextile/apparel 10–15%",
                "低十％案件が存在する程度",
                "selection bias大。TP主根拠にしない",
            ],
        ],
        col_widths=[1.8, 1.8, 2.6],
    )

    add_h(doc, "7.5 Retention仮説と感度", 2)
    add_para(
        doc,
        "関係式（概念）: retention% − Canada関連OPEX% ≈ 目標ROS%。",
        space_after=6,
    )
    add_table(
        doc,
        ["仮説retention（Net）", "位置づけ", "コメント"],
        [
            ["5%", "下側感度", "pure rep並み。SOR/与信ありでもOPEXが低ければROSは成立しうる。排除しない"],
            ["7.5%", "現状案件に近い候補", "既存/打診distributor管理が中心でMDが薄い場合に整合しやすい"],
            ["10%", "テスト仮説（中心候補の一つ）", "OPEXが売上比≈7–8%ならROS≈2–3%となり説明しやすい。未証OPEXのまま「実証済み」としない"],
            ["12.5%", "上側感度", "開拓費・travel・専任・重い与信が載る場合に合理化余地"],
            ["15%", "原則不採用", "Faire級のmarketplace/customer opsを担う米モデル寄りの率。独立CA卸相手では厚め"],
        ],
        col_widths=[1.4, 1.6, 3.2],
    )

    add_para(doc, "数値例（説明用）", bold=True, space_after=4)
    add_para(
        doc,
        "Canada Net Sales = 1,000,000 / retention 10% = 100,000 / OPEX 70,000 → EBIT 30,000 → ROS 3%。"
        "この3%が独立企業レンジに入るかを見る。OPEXが30,000しかないなら同一10%でROS 7%となり、routine仮説から外れるためretention引下げが必要。",
        size=10,
    )

    add_para(
        doc,
        "Gross 10%とNet 3–6%は「ベースが違うから実は近い」わけではない。"
        "直卸でGross≈Netなら Gross 10% ≒ Net 10% であり、Net 5%の約2倍になる。",
        size=10,
    )

    add_h(doc, "7.6 文書化の四段構成（修正版）", 2)
    add_numbered(
        doc,
        [
            "Industry: NAICS 424310 等でcomparable search universeを定義（GM37%は除外説明のみ）",
            "Supplemental: manufacturers’ rep commission 5–9%を下側参考。求人例は脚注",
            "FAR: TANAAKKはagentより厚く、full-inventory wholesaler/IORより薄い",
            "Primary test: retention適用後のCanada ROS/OMが外部レンジに入る（TNMM）",
        ],
    )

    add_h(doc, "7.7 精算設計（TPと会計の接続）", 2)
    add_bullets(
        doc,
        [
            "Net Sales = Distributor請求額 − 返品・値引 −（方針により）回収不能",
            "与信をTANAAKKが負うなら、未収のままLCJへ売上連動払いをしない",
            "国際輸送等の負担はScheduleで固定（配分前控除かLCJ owner費用か）",
            "四半期: 実績OMが2.0–3.5%バンドを外れたら retentionを契約バンド内で調整",
            "米15/85の機械適用禁止。刺しゅう/生地でCanada率が異なってもFAR優先",
        ],
    )

    add_h(doc, "7.8 TPワークプラン", 2)
    add_table(
        doc,
        ["Step", "作業", "成果物"],
        [
            ["T1", "FAR確定（本資料§3）とPLセグメント定義", "FARメモ"],
            ["T2", "Canada OPEX配賦ルール（人件費・旅費・与信・本社）", "配賦表"],
            ["T3", "retention感度（5/7.5/10/12.5）×ROS", "感度Excel"],
            ["T4", "DB等で近似comparables抽出（在庫薄・販売機能）", "セット一覧"],
            ["T5", "3–5年OM分布、必要ならWC/在庫強度調整", "レンジ"],
            ["T6", "契約バンドと調整条項をScheduleへ反映", "条文"],
            ["T7", "Local file / 社内TPメモ更新", "文書"],
        ],
        col_widths=[0.7, 3.0, 2.5],
    )

    # --- 8 Integrated roadmap ---
    add_h(doc, "8. 統合ロードマップ（時系列）", 1)

    add_table(
        doc,
        ["段階", "期間目安", "主担当", "主要タスク", "Exit criteria"],
        [
            [
                "Phase 0\n方針",
                "即時",
                "経営/TP/法務",
                "本方針承認。buy-sell必須論を採らない／No.2不流用／IOR=Distributor",
                "書面方針承認",
            ],
            [
                "Phase 1\n制度設計",
                "2–4週",
                "法務+税務+Broker",
                "Schedule骨子、Incoterms、broker質問状、GST/PE質問状、OPEX仮置き",
                "三者レビュー完了",
            ],
            [
                "Phase 2\n数値",
                "2–3週\n（並行可）",
                "財務/TP",
                "感度、仮retention（例: 7.5%と10%併記）、comparables着手",
                "仮説とバンド合意",
            ],
            [
                "Phase 3\n契約",
                "3–6週",
                "法務",
                "Canada Schedule署名、Distributor契約、オペSOP",
                "署名完了",
            ],
            [
                "Phase 4\nパイロット",
                "4–8週",
                "オペ+Broker",
                "初回PO〜通関〜入金〜精算のE2E",
                "不具合リスト解消",
            ],
            [
                "Phase 5\n定常",
                "継続",
                "全員",
                "四半期OMレビュー、GST再判定、率のtrue-up",
                "年次TP更新",
            ],
        ],
        col_widths=[1.0, 0.9, 1.0, 2.2, 1.2],
    )

    # --- 9 RACI risks ---
    add_h(doc, "9. リスク登録（要約）", 1)
    add_table(
        doc,
        ["リスク", "水準", "緩和"],
        [
            ["紙SOR化（LCJが実質値決め）", "高", "交渉禁止オペ、証跡、TANAAKK実機能"],
            ["米Path Bの沈黙輸入", "高", "Canada独自のtitle/精算定義"],
            ["書類ストーリー不一致", "高", "broker見本承認、SOP"],
            ["retention過大でOM過大", "中高", "OPEX計測、四半期調整、15%禁止"],
            ["GST登録漏れ/過度", "中", "規模シナリオで税理士確認、年次再判定"],
            ["与信と精算の不整合", "中", "回収ベース精算"],
            ["求人佣金への過度依存", "中", "TNMM主、佣金は補足"],
        ],
        col_widths=[2.2, 0.8, 3.2],
    )

    # --- 10 Decision ---
    add_h(doc, "10. 意思決定用サマリー", 1)
    add_para(doc, "採る", bold=True, space_after=4)
    add_bullets(
        doc,
        [
            "Japan直送、SOR=TANAAKK、IOR=Distributor、LCJ=在庫+fulfillment",
            "Consignment維持可（真正commercial sale + Canada Schedule再定義が条件）",
            "Canada専用契約、No.2不流用、15%不流用",
            "TPはTNMM（ROS）主、commissionは補足、10%はテスト仮説",
        ],
    )
    add_para(doc, "採らない", bold=True, space_after=4)
    add_bullets(
        doc,
        [
            "米経由在庫の再輸出前提",
            "LCJをCanada顧客の売主・交渉主体にする",
            "customsのためだけの形式的二段buy-sell必須化",
            "Gross 10%＝Net 3–6%と同値であるかのような説明",
            "外部データだけで8–12%/中心10%を「確認済み」と宣言すること",
        ],
    )

    add_h(doc, "11. 免責", 1)
    add_para(
        doc,
        "本資料は社内検討用の事業・TP方針整理であり、弁護士・税理士・通関士による個別事案の確定助言ではない。"
        "Canada税関評価、GST/HST、条約PE、日米移転価格文書は、それぞれ資格専門家の書面確認を経て実施すること。"
        "外部統計・MANA調査等は時点・定義に依存し、アームズレングスの唯一の証拠ではない。",
        size=10,
    )

    add_para(
        doc,
        "改訂履歴: 2026-10-07 初版（Canada distributor直送・FAR・契約/通関ロードマップ・TP戦略の議論統合）。",
        size=9,
        space_before=12,
    )

    OUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUT)
    print(f"Wrote {OUT}")


if __name__ == "__main__":
    build()
